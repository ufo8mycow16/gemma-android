"""Builds stand-in versions of Glen's Vulcan Fury documents for the regression tests.
The real script is not available here; these reproduce its structure and the traps
that broke v9 (a table of contents, incidental 'Episode 3 will reveal...' prose,
design sections mentioning episodes, adjacent episodes with distinctive content)."""
import io
from docx import Document

EP1 = ["Kirk, Spock and McCoy beam down to Vulcan for the ceremony.", "Chapel organises a bowling night on the recreation deck.",
       "The Science Officer's Log records the first anomaly."]
EP2 = ["Spock is summoned before the Vulcan High Council.", "McCoy argues with T'Pau about the ritual."]
EP3 = ["Kirk, Spock and McCoy remain on Vulcan while the Enterprise holds orbit.",
       "Romulan infiltrators hijack the Enterprise; the ship is severely undermanned.",
       "Scotty becomes acting captain and takes the conn.",
       "Control of the ship is traced to Auxiliary Control on Deck 7, where the Romulan Infiltration Matrix is found.",
       "Scotty drops the Enterprise into an asteroid field to shake the intruders' control.",
       "The damage to the Matrix is repaired and the transporter logs are reconstructed.",
       "The logs show the Romulans have already transported to the hidden mining colony.",
       "An away team attempts pursuit, but the transporter has been sabotaged.",
       "The intruders escape capture, and the trail leads on to Episode 4."]
EP4 = ["The away team reaches the hidden mining colony on Kalandra.", "Uhura decodes the Romulan transmissions."]
FILLER = "The corridor lights flicker as the crew moves between decks, and nobody speaks for a long moment. " * 6


def build(styled=True, with_toc=True, filler_paras=12):
    d = Document()
    def head(text, level):
        if styled:
            d.add_heading(text, level)
        else:
            d.add_paragraph(text)
    head("SECRET OF VULCAN FURY", 0)
    if with_toc:
        d.add_paragraph("Contents")
        for line in ("Episode 1 – Arrival", "Episode 2 – The Council", "Episode 3 – Acting Captain", "Episode 4 – The Colony"):
            d.add_paragraph(line)
    d.add_paragraph("Project notes: Episode 3 will reveal the full scale of the Romulan plan, so Episode 3 must not be cut.")
    for n, title, body in ((1, "Arrival", EP1), (2, "The Council", EP2), (3, "Acting Captain", EP3), (4, "The Colony", EP4)):
        head(f"Episode {n} – {title}", 1)
        for i, line in enumerate(body):
            d.add_paragraph(line)
            for _ in range(filler_paras // max(1, len(body))):
                d.add_paragraph(FILLER)
        if n == 2:
            d.add_paragraph("Episode 3 will reveal what the Council feared.")
    head("Design Bible: Character Portraits", 1)
    d.add_paragraph("Scotty's portrait is used in Episode 3 and in the HUD of the acting-captain sequence.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Episode", "Portrait"
    t.cell(1, 0).text, t.cell(1, 1).text = "Episode 3", "Scotty, command gold"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_episode1_master():
    d = Document()
    d.add_heading("Episode 1 – Arrival (Master)", 1)
    d.add_paragraph("MASTER VERSION: Kirk, Spock and McCoy arrive on Vulcan; T'Pring greets them at the gate.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
