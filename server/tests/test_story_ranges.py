"""
Story ranges: episodes found from their real extent, whatever the Word heading levels say.
Built to reproduce what the real Vulcan Fury master script does (October 2026 report):
  - front matter first ("Production Design Document Standard"), and the prologue's design documents
  - lists that name every episode (a variables table, portrait list) before any episode starts
  - Episode 1 as an unstyled title line, ended by "END OF EPISODE 1." and then an "EPISODE 1 COMPLETE" screen
  - Episode 2 with NO heading of its own, only text naming itself, and numbered sections up to "19."
  - Episode 3 headed at Heading 2 while its Acts are Heading 1 (inverted levels), ended by "Episode 3 ends."
  - a preview of Episode 4 between the end of Episode 3 and Episode 4's heading
  - Episode 4 opening with "Opening Recap — The End of Episode 3", ended by "END OF EPISODE 4."
  - Episode 5 headed several times
    python3 tests/test_story_ranges.py
"""
import base64, io, json, re, sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import docx  # noqa: E402
from gemma_core import documents as d, routing  # noqa: E402

checks = 0


def ok(cond, what):
    global checks
    checks += 1
    if not cond:
        raise AssertionError("FAILED: " + what)


def lorem(tag, n=12):
    return " ".join(f"{tag} event {i} happens here with the crew." for i in range(n))


def build() -> bytes:
    doc = docx.Document()
    h = lambda text, level: doc.add_heading(text, level)
    p = doc.add_paragraph
    h("SECRET OF VULCAN FURY", 1); h("Production Design Document Standard", 2)
    for i in range(1, 15):
        h(f"{i}. Scene Information Rule {i}", 1); p("Every scene document records its purpose and flow. " * 3)
    h("PROLOGUE", 3); p("Scene Name: Enterprise Recreation Deck – Bowling Alley. " + lorem("Prologue bowling", 20))
    h("Variable Categories", 1)
    for x in ("PROLOGUE", "Episode 1", "Episode 2", "Episode 3", "Episode 4", "Episode 5"):
        p(x)
    p("Global variables are shared across the entire game. " * 5)
    h("Prologue Variables", 1); p(lorem("Prologue variable", 10))
    for n, who in ((1, "McCoy"), (2, "Spock"), (3, "Scotty"), (4, "Chapel"), (5, "Kirk")):
        p(f"Episode {n}"); p(who)
    h("PROLOGUE RECORDING SCRIPT", 2); p(lorem("Prologue recording", 30))
    p("EPISODE 1 – THE DEAD DELEGATE")
    p("Episode 1 is Dr Leonard McCoy's episode. " + lorem("Episode one delegate", 40))
    h("1. EPISODE OVERVIEW", 1); p(lorem("Episode one overview", 30))
    p("This foreshadows Episode 2 and 3.")
    p("END OF EPISODE 1.")
    h("EPISODE 1 COMPLETE", 1); p("Score Categories. " + lorem("Episode one score", 8)); p("Episode 2 begins immediately from that line.")
    h("1. EPISODE OVERVIEW", 1)
    p("Episode 2 is Commander Spock's episode. Episode 2 investigates what remains inside his dead mind. " + lorem("Episode two meld", 30))
    for sec in ("6. EPISODE 2 PLAYER-KNOWLEDGE RULE", "12. SCENE 2.02 — MIND MELD ONE", "17.22 THE MELD BEGINS", "17.28 THE MELD ENDS",
                "19. EPISODE II FINALE"):
        h(sec, 1); p("Episode 2 continues. " + lorem(f"Episode two {sec[:5]}", 25))
    h("Episode 3 — Expanded Detailed Synopsis", 2)
    p("Episode 3 begins immediately after the Enterprise disappears from Vulcan orbit. Scotty takes command. " + lorem("Episode three scotty", 10))
    for act in ("Act One — The Hijacked Enterprise", "Act Two — Out of Control", "Act Eleven — The Deleted Transporter Record"):
        h(act, 1); p(lorem(f"Episode three {act[:8]}", 25))
    p("The deleted transporter records show the infiltrators beamed to the mining colony; the transporters are sabotaged.")
    p("Episode 3 ends.")
    p("Where Episode 4 Begins"); p("Episode 4 begins with the Enterprise transporter still unavailable. But that's Episode 5.")
    h("Episode 4 — Detailed Working Synopsis", 2)
    h("Opening Recap — The End of Episode 3", 3); p("Episode 4 fades in during the closing minutes of Episode 3: the deleted transporter records.")
    h("Episode 4 Begins", 1); p("The shuttle Galileo takes the away team to the mining colony. " + lorem("Episode four colony", 25))
    h("Act One — The Mining Colony", 1); p(lorem("Episode four mine", 25))
    p("END OF EPISODE 4.")
    p("Epiosde 5 Synopsis")
    h("Episode 5 — FURY", 2); p(lorem("Episode five fury", 25))
    h("ACT I — ENTERING FURY", 1); p(lorem("Episode five enter", 25))
    h("Episode 5 — FURY", 2); p(lorem("Episode five tri-meld", 25))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
data = build()
idx = d.build_index(d.extract_blocks(DOCX, data))
parts = [(r["kind"], r["number"], r["source"]) for r in idx.ranges]
ok(parts == [("prologue", None, "headings"), ("episode", 1, "headings"), ("episode", 2, "inferred"), ("episode", 3, "headings"),
             ("episode", 4, "headings"), ("episode", 5, "headings")], f"six story parts, Episode 2 inferred: {parts}")
by = {(r["kind"], r["number"]): r for r in idx.ranges}
t = lambda r: idx.text[r["start"]:r["end"]]
ok("Episode 1 is Dr Leonard McCoy" in t(by["episode", 1]) and "Score Categories" in t(by["episode", 1]), "Episode 1 runs through its COMPLETE screen")
ok("Episode 2 is Commander Spock" in t(by["episode", 2]) and "17.22 THE MELD BEGINS" in t(by["episode", 2]), "Episode 2 has its own text")
ok("Episode one" not in t(by["episode", 2]) and "Episode three" not in t(by["episode", 2]), "Episode 2 contains nothing of 1 or 3")
e3 = t(by["episode", 3])
ok("Act Eleven" in e3 and "deleted transporter records" in e3 and e3.rstrip().endswith("Episode 3 ends."),
   "Episode 3 includes its Acts (Heading 1 under a Heading 2) and ends at its end marker")
ok("shuttle Galileo" not in e3 and "Episode four" not in e3, "no Episode 4 text in Episode 3")
e4 = t(by["episode", 4])
ok(e4.startswith("Episode 4 — Detailed Working Synopsis") and "shuttle Galileo" in e4 and "Act One — The Mining Colony" in e4 and
   e4.rstrip().endswith("END OF EPISODE 4."), "Episode 4 runs from its heading to END OF EPISODE 4, across its Heading 1 parts")
ok("Episode five" not in e4, "no Episode 5 text in Episode 4")
ok("tri-meld" in t(by["episode", 5]), "Episode 5 continues across its repeated headings")
ok("Prologue bowling" in t(by["prologue", None]) and "Production Design Document Standard" not in t(by["prologue", None]),
   "the prologue starts at its own heading, after the front matter")
ok(not any(r["start"] < idx.text.find("EPISODE 1 – THE DEAD DELEGATE") for r in idx.ranges if r["kind"] == "episode"),
   "lists naming every episode are not episode starts")

# routing of Glen's four questions
qs = {"To make sure your programming is holding up, I'd like a brief synopsis of each episode please.": ("WHOLE_DOCUMENT_ANALYSIS", None),
      "is there an episode 3?": ("STRUCTURAL_DETAIL", ("episode", 3)),
      "tell me what happens in episode 4": ("STRUCTURAL_DETAIL", ("episode", 4)),
      "tell me what else happens after the intro": ("WHOLE_DOCUMENT_ANALYSIS", None)}
for q, (intent, target) in qs.items():
    r = routing.classify(q, documents_known=True)
    ok(r.intent == intent and r.target == target, f"route for {q!r}: {r.intent} {r.target}")

# end to end, through a real server: what Gemma is actually given for each question
from harness import Live  # noqa: E402
live = Live(scale="0.002")
try:
    cid = live.new_chat("story ranges")
    doc = {"name": "Master Script.docx", "mime": DOCX, "data": f"data:{DOCX};base64," + base64.b64encode(data).decode()}
    live.turn(cid, "Here's the complete master script.", documents=[doc], timeout=600)

    def given(q):
        n = len(live.chat_requests())
        live.turn(cid, q, timeout=600)
        req = [r for r in live.chat_requests()[n:] if r["purpose"] == "chat"][-1]
        return "\n".join(m["content"] for m in req["messages"] if m["role"] in ("system", "user"))

    g = given("To make sure your programming is holding up, I'd like a brief synopsis of each episode please.")
    ok("<document_map" in g and "This is the complete uploaded document" in g, "the whole-script answer is told what the document is")
    ok(len(re.findall(r"<story_part ", g)) == 6 and 'part="Episode 2"' in g, "each of the six parts is given separately, Episode 2 included")
    ok("<document_outline" not in g, "no front-matter outline standing in for the story")
    ok("never extra episodes" in g and "(5 in all)" in g, "she is told how many episodes there are")
    g = given("tell me what happens in episode 4")
    sec = re.search(r"<document_section [^>]*>\n(.*?)</document_section>", g, re.S).group(1)
    ok("shuttle Galileo" in sec and "Episode three" not in sec and "Episode five" not in sec, "Episode 4 evidence is Episode 4's range only")
    ok("recaps or refers to Episode 3" in sec, "the opening recap is marked as Episode 3's events")
    g = given("is there an episode 3?")
    sec = re.search(r"<document_section [^>]*>\n(.*?)</document_section>", g, re.S).group(1)
    ok("Act Eleven" in sec and "shuttle Galileo" not in sec, "Episode 3 evidence includes its acts and stops before Episode 4")
    g = given("tell me what else happens after the intro")
    ok(len(re.findall(r"<story_part ", g)) == 6, "'what else happens after the intro' gets the whole story, part by part")
    g = given("what's in episode 2?")
    ok('section="Episode 2"' in g and "Episode 2 is Commander Spock" in g, "Episode 2 is found although it has no heading")
finally:
    live.stop()
print(f"STORY RANGE TESTS OK: {checks} checks")
